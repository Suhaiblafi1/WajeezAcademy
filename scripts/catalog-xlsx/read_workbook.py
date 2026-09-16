#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
قارئُ كرّاسةِ المراجعةِ العائدة — **تقريرٌ لا كتابة**.

لا يمسُّ هذا السكربتُ ملفَّ الكتالوجِ إطلاقا. يقرأ الأعمدةَ الصفراءَ وحدَها،
ويقابلها بمصادرِ المستودعِ كما هي الآن، ويُخرج قائمةَ الفروقِ لتُراجَع قبل أن
يُكتب شيء. والكتابةُ خطوةٌ تالية، بعد أن يرى صاحبُ المنصّةُ ما سيتغيّر.

ولمَ لا يُقرأ العمودُ الرمادي؟ لأنّ المرجعَ هو المستودعُ لا الكرّاسة. فلو مرّ
على الملفِّ شهرٌ وتغيّر الكتالوجُ في أثنائه، حُسب الفرقُ على الحاضرِ لا على
صورةٍ قديمةٍ محفوظةٍ في الملفّ — ولو عبث أحدٌ بعمودِ الحالي لم يُفسد شيئا.

    python3 scripts/catalog-xlsx/read_workbook.py العائد.xlsx [--json تقرير.json]
"""

from __future__ import annotations

import json
import os
import sys
from typing import Any

from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
CATALOG = os.path.join(ROOT, "src", "data", "catalog", "core-catalog.v2.json")
SKILLS = os.path.join(ROOT, "src", "data", "catalog", "skills.v1.ar.json")

SENTINEL = "(حذف)"
EDIT_MARK = "⟵"

# حدودُ بوّابةِ المصدر (server/catalog/validate-source.ts) — تُفحص هنا مبكّرا
HOURS_MIN, HOURS_MAX = 1, 40
PRICE_MIN, PRICE_MAX = 100, 200

LIST_FIELDS = {
    "skill_slugs", "skill_ids", "skill_names_ar", "learning_outcomes_ar",
    "learning_objectives_ar", "rubric_criteria_ar", "source_codes",
}
INT_FIELDS = {
    "total_hours", "guided_hours", "independent_hours", "practice_hours",
    "sequence", "expected_hours", "list_price", "suggested_total_hours",
}
# حقولٌ تخاطبني ولا تدخل الكتالوجَ حقلا
META_FIELDS = {"_action", "_notes", "_decision"}

SHEETS = {
    "الدورات": ("course_id", "courses"),
    "المحاور": ("module_id", "modules"),
    "محتوى_الدروس": ("module_id", "modules"),
    "المهارات": ("skill_id", "skills"),
    "مهارات_بلا_دورة": ("skill_id", "skills"),
    "دورات_مقترحة": ("proposed_id", "proposed"),
}


def norm(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).replace("\r\n", "\n").strip()


def load_sources() -> dict[str, dict]:
    with open(CATALOG, encoding="utf-8") as f:
        cat = json.load(f)
    with open(SKILLS, encoding="utf-8") as f:
        sk = json.load(f)
    idx = {
        "courses": {c["course_id"]: c for c in cat["courses"]},
        "modules": {m["module_id"]: m for m in cat["modules"]},
        "skills": {s["skill_id"]: s for s in sk["skills"]},
        "proposed": {},
    }
    known = {s["slug"] for s in sk["skills"]}
    known |= {e.get("slug") for e in cat.get("skill_extensions", []) if e.get("slug")}
    idx["_known_slugs"] = known
    idx["_currencies"] = {c.get("list_currency", "USD") for c in cat["courses"]}
    return idx


def current_value(rec: dict, field: str) -> str:
    v = rec.get(field)
    if v is None:
        return ""
    if isinstance(v, list):
        return "\n".join(str(x) for x in v)
    if isinstance(v, bool):
        return "نعم" if v else "لا"
    return str(v)


def validate(field: str, value: str, idx: dict) -> str | None:
    """يُرجع سببَ الرفضِ أو None. يُفحص مبكّرا لئلّا يُكتشف في CI بعد الالتزام."""
    if value == "":
        return None
    if field in INT_FIELDS:
        try:
            n = int(float(value))
        except ValueError:
            return f"«{value}» ليس عددا صحيحا"
        if field == "total_hours" and not (HOURS_MIN <= n <= HOURS_MAX):
            return f"الساعات {n} خارج المدى [{HOURS_MIN}، {HOURS_MAX}] — ترفضه البوّابة"
        if field == "list_price" and not (PRICE_MIN <= n <= PRICE_MAX):
            return f"السعر {n} خارج المدى المعتمد [{PRICE_MIN}، {PRICE_MAX}] — ترفضه البوّابة"
    if field == "skill_slugs":
        unknown = [s.strip() for s in value.split("\n") if s.strip() and s.strip() not in idx["_known_slugs"]]
        if unknown:
            return "مهاراتٌ غير مسجَّلة: " + " · ".join(unknown) + " — تُسجَّل أوّلا وإلّا رُفض الاستيراد"
    if field == "list_currency":
        others = idx["_currencies"] - {value}
        if others:
            return f"عملةٌ تخالف بقيّةَ الكتالوج ({' · '.join(sorted(others))}) — لا يُجمع مسارٌ بعملتين"
    return None


def main() -> int:
    if len(sys.argv) < 2:
        print(__doc__)
        return 2
    path = sys.argv[1]
    if not os.path.exists(path):
        print(f"✗ لا ملفَّ عند {path}")
        return 2

    idx = load_sources()
    wb = load_workbook(path, data_only=True)

    changes: list[dict] = []
    problems: list[dict] = []
    notes: list[dict] = []
    new_rows: list[dict] = []
    noops = 0

    for sheet, (key_field, bucket) in SHEETS.items():
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]

        # عمودُ التعديلِ يُعرف من ترويسته، والحقلُ من صفِّ المفاتيحِ يسارَه
        edits: list[tuple[str, int]] = []
        for c in range(1, ws.max_column + 1):
            if str(ws.cell(row=2, column=c).value or "").startswith(EDIT_MARK):
                field = ws.cell(row=3, column=c - 1).value
                if field:
                    edits.append((str(field), c))

        for r in range(4, ws.max_row + 1):
            rid = norm(ws.cell(row=r, column=1).value)
            row_vals = {f: norm(ws.cell(row=r, column=c).value) for f, c in edits}
            if not rid:
                if any(row_vals.values()):
                    new_rows.append({"sheet": sheet, "row": r, "values": {k: v for k, v in row_vals.items() if v}})
                continue

            rec = idx.get(bucket, {}).get(rid)
            for field, val in row_vals.items():
                if val == "":
                    continue
                if field in META_FIELDS:
                    notes.append({"sheet": sheet, "id": rid, "field": field, "value": val})
                    continue
                if rec is None:
                    problems.append({"sheet": sheet, "id": rid, "field": field,
                                     "reason": "معرّفٌ لا أصلَ له في المصدر"})
                    continue
                cur = current_value(rec, field)
                new = "" if val == SENTINEL else val
                if norm(cur) == norm(new):
                    noops += 1
                    continue
                why = validate(field, new, idx)
                entry = {"sheet": sheet, "id": rid, "field": field,
                         "from": cur, "to": new, "clear": val == SENTINEL}
                if why:
                    entry["blocked"] = why
                    problems.append(entry)
                else:
                    changes.append(entry)

    # ── التقرير ───────────────────────────────────────────────────
    def trim(s: str, n: int = 70) -> str:
        s = s.replace("\n", " ⏎ ")
        return s if len(s) <= n else s[: n - 1] + "…"

    print("═" * 78)
    print(f"تقريرُ الفروق · {os.path.basename(path)}  —  قراءةٌ فقط، لم يُكتب شيء")
    print("═" * 78)
    print(f"  تغييراتٌ مقبولة : {len(changes)}")
    print(f"  تردّها البوّابة  : {len(problems)}")
    print(f"  صفوفٌ جديدة     : {len(new_rows)}")
    print(f"  ملاحظاتٌ وقرارات: {len(notes)}")
    print(f"  بلا أثر (مطابق) : {noops}")

    if changes:
        print("\n── تغييراتٌ ستُطبَّق " + "─" * 57)
        for ch in changes:
            head = f"  [{ch['sheet']}] {ch['id']} · {ch['field']}"
            print(head + ("   ← تفريغ" if ch["clear"] else ""))
            print(f"      من: {trim(ch['from'])}")
            print(f"      إلى: {trim(ch['to'])}")

    if problems:
        print("\n── تردّها البوّابة — لا تُطبَّق حتّى تُحلّ " + "─" * 35)
        for p in problems:
            print(f"  [{p['sheet']}] {p['id']} · {p.get('field','')}")
            print(f"      ✗ {p.get('blocked') or p.get('reason')}")

    if new_rows:
        print("\n── صفوفٌ جديدةٌ بلا معرّف " + "─" * 52)
        for n in new_rows:
            print(f"  [{n['sheet']}] صفّ {n['row']}: {trim(json.dumps(n['values'], ensure_ascii=False), 90)}")

    if notes:
        print("\n── قراراتٌ وملاحظاتٌ لي " + "─" * 54)
        for n in notes:
            print(f"  [{n['sheet']}] {n['id']} · {n['field']}: {trim(n['value'])}")

    print("\n" + "═" * 78)
    if problems:
        print("لا يُكتب شيءٌ ما بقي مردودٌ بلا حلّ — تُراجَع أوّلا.")
    elif changes or new_rows or notes:
        print("جاهزٌ للتطبيق بعد موافقتك على القائمة أعلاه.")
    else:
        print("لا فرقَ — الكرّاسةُ مطابقةٌ للمصدر.")

    if "--json" in sys.argv:
        out = sys.argv[sys.argv.index("--json") + 1]
        with open(out, "w", encoding="utf-8") as f:
            json.dump({"changes": changes, "problems": problems,
                       "new_rows": new_rows, "notes": notes, "noops": noops},
                      f, ensure_ascii=False, indent=1)
        print(f"وكُتب التقريرُ إلى {out}")

    return 1 if problems else 0


if __name__ == "__main__":
    sys.exit(main())
