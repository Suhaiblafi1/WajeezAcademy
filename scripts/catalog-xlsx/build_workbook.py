#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
مولّدُ كرّاسةِ مراجعةِ الكتالوج (Excel).

يقرأ مصادرَ الكتالوجِ كما هي في الشجرة، ويُخرج ملفّا واحدا فيه لكلِّ حقلٍ
عمودان: «الحالي» كما في المصدر، و«تعديلك» فارغٌ يكتب فيه صاحبُ المنصّة.

والقاعدةُ التي عليها الرحلةُ كلُّها: **الفارغُ لا يعني الحذف، بل «لا تغيير»**.
ومن أراد إفراغَ حقلٍ كتب `(حذف)` صريحةً — وبلا هذا التمييزِ لا يُفرَّق بين
حقلٍ لم يُمسّ وحقلٍ أُريد محوُه، فتُمحى ثمانون دورةً بصمت.

ولا يُقرأ عمودُ «الحالي» عند العودة إطلاقا: الفرقُ يُحسب من ملفّاتِ المصدر
نفسِها لا من الكرّاسة، فلو عبث أحدٌ بعمودِ الحالي لم يُفسد الاستيراد.

    python3 scripts/catalog-xlsx/build_workbook.py [الوجهة.xlsx]
"""

from __future__ import annotations

import json
import os
import re
import sys
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
CATALOG = os.path.join(ROOT, "src", "data", "catalog", "core-catalog.v2.json")
SKILLS = os.path.join(ROOT, "src", "data", "catalog", "skills.v1.ar.json")
PROPOSED = os.path.join(os.path.dirname(os.path.abspath(__file__)), "proposed-courses.json")

FONT = "Arial"
SENTINEL = "(حذف)"

# ألوان: الترويسةُ داكنة، والحاليُّ رماديٌّ لا يُكتب فيه، والتعديلُ أصفرُ يُكتب فيه
C_HEAD = "1F4E5F"
C_HEAD_EDIT = "8A6D0B"
C_KEY = "DCE9F0"
C_NOW = "F2F2F2"
C_EDIT = "FFF7D6"
C_NOTE = "EAF4EA"
C_WARN = "FDE9E9"

ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")

thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def clean(v: Any) -> Any:
    """يُسقط المحارفَ التي يرفضها XLSX، ويضمّ القوائمَ سطرا لكلِّ عنصر."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "نعم" if v else "لا"
    if isinstance(v, (int, float)):
        return v
    if isinstance(v, list):
        return "\n".join(str(clean(x)) for x in v)
    if isinstance(v, dict):
        return json.dumps(v, ensure_ascii=False)
    return ILLEGAL.sub("", str(v))


class Sheet:
    """ورقةٌ أعمدتُها أزواج: حاليٌّ يُقرأ، وتعديلٌ يُكتب."""

    def __init__(self, wb: Workbook, name: str, intro: str):
        self.ws = wb.create_sheet(name)
        self.ws.sheet_view.rightToLeft = True
        self.intro = intro
        self.cols: list[tuple[str, str, str, int]] = []  # (مفتاح، عنوان، نوع، عرض)

    def col(self, key: str, label: str, kind: str = "edit", width: int = 30):
        """kind: key = مفتاحٌ لا يُعدَّل · edit = حقلٌ له عمودُ تعديل · note = عمودُ ملاحظة"""
        self.cols.append((key, label, kind, width))

    def build(self, rows: list[dict]):
        ws = self.ws
        ws["A1"] = self.intro
        ws["A1"].font = Font(name=FONT, size=11, bold=True, color="1F4E5F")
        ws["A1"].alignment = Alignment(horizontal="right", vertical="center")
        ws.row_dimensions[1].height = 30

        c = 1
        pos: list[tuple[str, str, int]] = []  # (مفتاح، نوع، عمود)
        for key, label, kind, width in self.cols:
            head = ws.cell(row=2, column=c, value=label)
            head.fill = PatternFill("solid", fgColor=C_HEAD)
            head.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
            head.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            head.border = BORDER
            sub = ws.cell(row=3, column=c, value=key)
            sub.fill = PatternFill("solid", fgColor=C_HEAD)
            sub.font = Font(name=FONT, size=8, italic=True, color="BFD7E2")
            sub.alignment = Alignment(horizontal="center", vertical="center")
            sub.border = BORDER
            ws.column_dimensions[get_column_letter(c)].width = width
            pos.append((key, kind, c))
            c += 1

            if kind == "edit":
                eh = ws.cell(row=2, column=c, value=f"⟵ تعديلك · {label}")
                eh.fill = PatternFill("solid", fgColor=C_HEAD_EDIT)
                eh.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
                eh.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                eh.border = BORDER
                es = ws.cell(row=3, column=c, value="اتركه فارغا = لا تغيير")
                es.fill = PatternFill("solid", fgColor=C_HEAD_EDIT)
                es.font = Font(name=FONT, size=8, italic=True, color="F2E2B0")
                es.alignment = Alignment(horizontal="center", vertical="center")
                es.border = BORDER
                ws.column_dimensions[get_column_letter(c)].width = width
                c += 1

        ws.row_dimensions[2].height = 42
        ws.row_dimensions[3].height = 16

        for i, row in enumerate(rows):
            r = 4 + i
            for key, kind, col in pos:
                cell = ws.cell(row=r, column=col, value=clean(row.get(key)))
                cell.font = Font(name=FONT, size=10)
                cell.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
                cell.border = BORDER
                if kind == "key":
                    cell.fill = PatternFill("solid", fgColor=C_KEY)
                    cell.font = Font(name=FONT, size=10, bold=True)
                elif kind == "note":
                    cell.fill = PatternFill("solid", fgColor=C_NOTE)
                else:
                    cell.fill = PatternFill("solid", fgColor=C_NOW)
                if kind == "edit":
                    e = ws.cell(row=r, column=col + 1)
                    e.fill = PatternFill("solid", fgColor=C_EDIT)
                    e.font = Font(name=FONT, size=10)
                    e.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
                    e.border = BORDER
            ws.row_dimensions[r].height = 58

        # أوّلُ عمودٍ وترويستُه ثابتان عند التمرير
        ws.freeze_panes = ws.cell(row=4, column=2)
        ws.auto_filter.ref = f"A2:{get_column_letter(c - 1)}{3 + len(rows)}"
        return self


def dropdown(sheet: Sheet, key: str, options: list[str], n_rows: int):
    """قائمةٌ منسدلةٌ على عمودِ التعديلِ وحدَه — لا على عمودِ الحالي."""
    col = 1
    for k, _label, kind, _w in sheet.cols:
        if k == key and kind == "edit":
            letter = get_column_letter(col + 1)
            dv = DataValidation(
                type="list",
                formula1='"' + ",".join(options) + '"',
                allow_blank=True,
                showDropDown=False,
            )
            sheet.ws.add_data_validation(dv)
            dv.add(f"{letter}4:{letter}{3 + n_rows}")
            return
        col += 2 if kind == "edit" else 1


def main() -> int:
    out = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, "دورات-وجيز-للمراجعة.xlsx")

    with open(CATALOG, encoding="utf-8") as f:
        cat = json.load(f)
    with open(SKILLS, encoding="utf-8") as f:
        sk = json.load(f)
    with open(PROPOSED, encoding="utf-8") as f:
        prop = json.load(f)

    courses = cat["courses"]
    modules = cat["modules"]
    pathways = cat["launch_pathways"]
    registry = sk["skills"]
    extensions = cat.get("skill_extensions", [])

    by_course = {c["course_id"]: c for c in courses}
    pw_title = {p["id"]: p.get("short_title") or p.get("title", "") for p in pathways}

    linked: dict[str, list[str]] = {}
    for c in courses:
        for s in c.get("skill_slugs", []) or []:
            linked.setdefault(s, []).append(c["course_id"])

    wb = Workbook()
    wb.remove(wb.active)

    # ═══ ١ · الدورات ═══════════════════════════════════════════════
    sh = Sheet(wb, "الدورات", "الدورات الحاليّة · ٨١ دورة. الرماديُّ هو الحاليُّ ولا يُقرأ عند العودة، والأصفرُ هو موضعُ تعديلك. والفارغُ = لا تغيير.")
    sh.col("course_id", "معرّف الدورة", "key", 16)
    sh.col("pathway_title", "المسار", "note", 26)
    sh.col("pathway_id", "معرّف المسار", "edit", 14)
    sh.col("sequence", "الترتيب في المسار", "edit", 10)
    sh.col("title_ar", "اسم الدورة", "edit", 38)
    sh.col("title_term_en", "المصطلح بالإنجليزية", "edit", 28)
    sh.col("subtitle_ar", "العنوان الفرعي", "edit", 42)
    sh.col("level_ar", "المستوى", "edit", 14)
    sh.col("total_hours", "مجموع الساعات", "edit", 12)
    sh.col("guided_hours", "ساعات موجَّهة", "edit", 12)
    sh.col("independent_hours", "ساعات ذاتية", "edit", 12)
    sh.col("practice_hours", "ساعات تطبيق", "edit", 12)
    sh.col("module_count", "عدد المحاور", "note", 10)
    sh.col("list_price", "سعر القائمة", "edit", 12)
    sh.col("list_currency", "العملة", "edit", 10)
    sh.col("short_promise_ar", "الوعد المختصر", "edit", 40)
    sh.col("description_ar", "الوصف", "edit", 50)
    sh.col("target_audience_ar", "الفئة المستهدفة", "edit", 40)
    sh.col("prerequisites_ar", "المتطلّبات السابقة", "edit", 34)
    sh.col("learning_objectives_ar", "أهداف التعلّم", "edit", 50)
    sh.col("learning_outcomes_ar", "مخرجات التعلّم", "edit", 50)
    sh.col("skill_count", "عدد المهارات", "note", 10)
    sh.col("skill_slugs", "معرّفات المهارات", "edit", 30)
    sh.col("skill_names_ar", "أسماء المهارات", "edit", 34)
    sh.col("practical_project_ar", "المشروع العملي", "edit", 46)
    sh.col("summative_assessment_ar", "التقييم الختامي", "edit", 46)
    sh.col("assessment_claim_ar", "دعوى التقييم", "edit", 44)
    sh.col("evidence_required_ar", "الدليل المطلوب", "edit", 40)
    sh.col("rubric_criteria_ar", "معايير الحكم", "edit", 30)
    sh.col("pass_rule_ar", "قاعدة النجاح", "edit", 30)
    sh.col("instructor_profile_ar", "مواصفات المدرّب", "edit", 40)
    sh.col("ai_policy_ar", "سياسة الذكاء الاصطناعي", "edit", 40)
    sh.col("legacy_title_ar", "الاسم القديم", "edit", 30)
    sh.col("source_codes", "رموز المصادر", "edit", 16)
    sh.col("standalone", "دورة مستقلّة", "edit", 12)
    sh.col("standalone_reason_ar", "سبب الاستقلال", "edit", 30)
    sh.col("_action", "إجراء مطلوب", "edit", 16)
    sh.col("_notes", "ملاحظاتك الحرّة", "edit", 40)

    mod_count: dict[str, int] = {}
    for m in modules:
        mod_count[m["course_id"]] = mod_count.get(m["course_id"], 0) + 1

    rows = []
    for c in sorted(courses, key=lambda x: (x.get("pathway_id") or "", x.get("sequence") or 0)):
        r = dict(c)
        r["pathway_title"] = pw_title.get(c.get("pathway_id"), "")
        r["module_count"] = mod_count.get(c["course_id"], 0)
        r["skill_count"] = len(c.get("skill_slugs") or [])
        rows.append(r)
    sh.build(rows)
    dropdown(sh, "level_ar", ["تأسيسي", "متوسط", "متقدّم"], len(rows))
    dropdown(sh, "_action", ["إبقاء", "أرشفة", "حذف", "مراجعة لاحقا"], len(rows))
    dropdown(sh, "list_currency", ["USD", "SAR", "AED"], len(rows))
    n_courses = len(rows)

    # ═══ ٢ · المحاور ═══════════════════════════════════════════════
    sh = Sheet(wb, "المحاور", "محاور الدورات · ٤٠٤ محورا. عنوانُ المحورِ ومخرجُه ونشاطُه وساعاتُه. ونصُّ الدرسِ نفسُه في ورقة «محتوى_الدروس».")
    sh.col("module_id", "معرّف المحور", "key", 18)
    sh.col("course_id", "الدورة", "note", 14)
    sh.col("course_title", "اسم الدورة", "note", 32)
    sh.col("sequence", "الترتيب", "edit", 9)
    sh.col("title_ar", "عنوان المحور", "edit", 38)
    sh.col("module_outcome_ar", "مخرج المحور", "edit", 44)
    sh.col("practice_activity_ar", "نشاط التطبيق", "edit", 40)
    sh.col("evidence_artifact_ar", "أثر التطبيق", "edit", 40)
    sh.col("expected_hours", "الساعات المتوقّعة", "edit", 12)
    sh.col("has_body", "له نصّ درس؟", "note", 12)
    sh.col("_action", "إجراء مطلوب", "edit", 16)
    sh.col("_notes", "ملاحظاتك الحرّة", "edit", 40)

    mrows = []
    for m in sorted(modules, key=lambda x: (x["course_id"], x.get("sequence") or 0)):
        r = dict(m)
        r["course_title"] = (by_course.get(m["course_id"]) or {}).get("title_ar", "")
        r["has_body"] = "نعم" if (m.get("module_body_ar") or "").strip() else "لا"
        mrows.append(r)
    sh.build(mrows)
    dropdown(sh, "_action", ["إبقاء", "أرشفة", "حذف", "مراجعة لاحقا"], len(mrows))
    n_modules = len(mrows)

    # ═══ ٣ · محتوى الدروس ══════════════════════════════════════════
    sh = Sheet(wb, "محتوى_الدروس", "نصوص الدروس المكتوبة · المحاور التي لها متنٌ اليوم. الصيغةُ Markdown، وتُعاد كما هي.")
    sh.col("module_id", "معرّف المحور", "key", 18)
    sh.col("course_title", "الدورة", "note", 28)
    sh.col("title_ar", "عنوان المحور", "note", 32)
    sh.col("module_body_ar", "متن الدرس (Markdown)", "edit", 80)
    sh.col("module_checks_ar", "تمرين الاسترجاع", "edit", 50)
    sh.col("module_scenario_ar", "السيناريو", "edit", 50)
    sh.col("module_practice_ar", "التطبيق", "edit", 46)
    sh.col("module_rubric_ar", "معيار الحكم", "edit", 46)
    sh.col("_notes", "ملاحظاتك الحرّة", "edit", 36)

    crows = []
    for m in sorted(modules, key=lambda x: (x["course_id"], x.get("sequence") or 0)):
        if not (m.get("module_body_ar") or "").strip():
            continue
        r = dict(m)
        r["course_title"] = (by_course.get(m["course_id"]) or {}).get("title_ar", "")
        crows.append(r)
    sh.build(crows)
    n_content = len(crows)

    # ═══ ٤ · المهارات ══════════════════════════════════════════════
    sh = Sheet(wb, "المهارات", "كلّ المهارات · سجلّ المهارات والامتدادات معا، ومع كلِّ مهارةٍ الدوراتُ التي تخدمها.")
    sh.col("skill_id", "معرّف المهارة", "key", 16)
    sh.col("slug", "المعرّف اللفظي", "note", 26)
    sh.col("name_ar", "اسم المهارة", "edit", 30)
    sh.col("name_en", "الاسم بالإنجليزية", "edit", 28)
    sh.col("family_ar", "العائلة", "note", 28)
    sh.col("family_id", "رمز العائلة", "note", 10)
    sh.col("level_band", "الطبقة", "edit", 14)
    sh.col("skill_type", "النوع", "edit", 12)
    sh.col("definition_ar", "التعريف", "edit", 50)
    sh.col("why_it_matters_ar", "لماذا تهمّ", "edit", 46)
    sh.col("source", "المصدر", "note", 16)
    sh.col("courses_count", "عدد الدورات", "note", 10)
    sh.col("courses", "الدورات التي تخدمها", "note", 34)
    sh.col("_notes", "ملاحظاتك الحرّة", "edit", 36)

    srows = []
    for s in registry:
        r = dict(s)
        cs = linked.get(s["slug"], [])
        r["courses_count"] = len(cs)
        r["courses"] = "\n".join(cs)
        srows.append(r)
    for e in extensions:
        r = dict(e)
        r["source"] = "skill_extensions (داخل الكتالوج)"
        cs = linked.get(e.get("slug", ""), [])
        r["courses_count"] = len(cs)
        r["courses"] = "\n".join(cs)
        srows.append(r)
    srows.sort(key=lambda x: (x.get("family_id") or "", x.get("skill_id") or ""))
    sh.build(srows)
    n_skills = len(srows)

    # ═══ ٥ · مهارات بلا دورة ═══════════════════════════════════════
    sh = Sheet(wb, "مهارات_بلا_دورة", "المهارات التي لا تخدمها دورةٌ اليوم · ٧٣ مهارة. ومع كلٍّ منها الدورةُ المقترحةُ في ورقة «دورات_مقترحة».")
    sh.col("skill_id", "معرّف المهارة", "key", 16)
    sh.col("slug", "المعرّف اللفظي", "note", 26)
    sh.col("name_ar", "اسم المهارة", "note", 30)
    sh.col("name_en", "الاسم بالإنجليزية", "note", 28)
    sh.col("family_ar", "العائلة", "note", 28)
    sh.col("level_band", "الطبقة", "note", 14)
    sh.col("definition_ar", "التعريف", "note", 50)
    sh.col("why_it_matters_ar", "لماذا تهمّ", "note", 46)
    sh.col("proposed_course", "الدورة المقترحة (اقتراحي)", "edit", 40)
    sh.col("_decision", "قرارك", "edit", 18)
    sh.col("_notes", "ملاحظاتك الحرّة", "edit", 40)

    slug_to_prop = {}
    for pc in prop["courses"]:
        for s in pc["skill_slugs"]:
            slug_to_prop[s] = f'{pc["proposed_id"]} — {pc["title_ar"]}'

    orows = []
    for s in registry:
        if s["slug"] in linked or s.get("active") is False or s.get("merged_into"):
            continue
        r = dict(s)
        r["proposed_course"] = slug_to_prop.get(s["slug"], "")
        orows.append(r)
    orows.sort(key=lambda x: (x.get("family_id") or "", x.get("skill_id") or ""))
    sh.build(orows)
    dropdown(sh, "_decision", ["أوافق على المقترح", "دورة أخرى", "دمج في دورة قائمة", "تُلغى المهارة", "تأجيل"], len(orows))
    n_orphans = len(orows)

    # ═══ ٦ · دورات مقترحة ══════════════════════════════════════════
    sh = Sheet(wb, "دورات_مقترحة", "دوراتٌ مقترحةٌ تغطّي المهاراتِ المعطّلة · اقتراحٌ لا قرار. لا يدخل منها الكتالوجَ شيءٌ حتّى تعتمده.")
    sh.col("proposed_id", "معرّف مقترح", "key", 16)
    sh.col("title_ar", "اسم الدورة المقترح", "edit", 44)
    sh.col("title_term_en", "المصطلح بالإنجليزية", "edit", 30)
    sh.col("family_id", "العائلة", "note", 12)
    sh.col("level_ar", "المستوى", "edit", 14)
    sh.col("suggested_total_hours", "الساعات المقترحة", "edit", 12)
    sh.col("skill_count", "عدد المهارات", "note", 10)
    sh.col("skill_slugs", "المهارات التي تغطّيها", "edit", 34)
    sh.col("skill_names_ar", "أسماء المهارات", "note", 36)
    sh.col("recommendation_ar", "توصيتي", "note", 24)
    sh.col("rationale_ar", "لماذا هذا التجميع", "note", 60)
    sh.col("_decision", "قرارك", "edit", 18)
    sh.col("_notes", "ملاحظاتك الحرّة", "edit", 44)

    name_of = {s["slug"]: s.get("name_ar", "") for s in registry}
    prows = []
    for pc in prop["courses"]:
        r = dict(pc)
        r["skill_count"] = len(pc["skill_slugs"])
        r["skill_names_ar"] = "\n".join(name_of.get(s, s) for s in pc["skill_slugs"])
        r.setdefault("recommendation_ar", "تُنشأ كما هي")
        prows.append(r)
    sh.build(prows)
    dropdown(sh, "_decision", ["أوافق", "أوافق مع تعديل", "أرفض", "أجّل"], len(prows))
    n_prop = len(prows)

    # ═══ ٧ · المسارات ══════════════════════════════════════════════
    sh = Sheet(wb, "المسارات", "المسارات · ٢٠ مسارا — للمرجع عند نقل دورةٍ من مسارٍ إلى آخر.")
    sh.col("id", "معرّف المسار", "key", 16)
    sh.col("short_title", "الاسم المختصر", "note", 30)
    sh.col("title", "الاسم الكامل", "note", 50)
    sh.col("audience", "الفئة", "note", 50)
    sh.col("course_count", "عدد الدورات", "note", 10)
    sh.col("courses", "دوراته", "note", 40)

    pw_rows = []
    for p in pathways:
        cs = [c["course_id"] for c in courses if c.get("pathway_id") == p["id"]]
        r = dict(p)
        r["course_count"] = len(cs)
        r["courses"] = "\n".join(cs)
        pw_rows.append(r)
    sh.build(pw_rows)

    # ═══ ٠ · دليل الاستخدام ════════════════════════════════════════
    g = wb.create_sheet("دليل_الاستخدام", 0)
    g.sheet_view.rightToLeft = True
    g.column_dimensions["A"].width = 34
    g.column_dimensions["B"].width = 96

    def line(r: int, a: str, b: str = "", head: bool = False, warn: bool = False):
        ca, cb = g.cell(row=r, column=1, value=a), g.cell(row=r, column=2, value=b)
        if head:
            ca.fill = PatternFill("solid", fgColor=C_HEAD)
            cb.fill = PatternFill("solid", fgColor=C_HEAD)
            ca.font = Font(name=FONT, size=12, bold=True, color="FFFFFF")
            cb.font = Font(name=FONT, size=12, bold=True, color="FFFFFF")
        else:
            ca.font = Font(name=FONT, size=10, bold=True)
            cb.font = Font(name=FONT, size=10)
            if warn:
                ca.fill = PatternFill("solid", fgColor=C_WARN)
                cb.fill = PatternFill("solid", fgColor=C_WARN)
        ca.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
        cb.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
        g.row_dimensions[r].height = 30 if not head else 26

    r = 1
    line(r, "كرّاسة مراجعة كتالوج وجيز", "ملفٌّ واحدٌ يخرج من المستودع ويعود إليه", head=True); r += 2

    line(r, "القاعدة الأولى", "لكلِّ حقلٍ عمودان: رماديٌّ فيه الحاليُّ، وأصفرُ بجانبه تكتب فيه الجديد. لا تكتب في الرمادي — فهو لا يُقرأ أصلا عند العودة."); r += 1
    line(r, "القاعدة الثانية", "العمودُ الأصفرُ الفارغ = لا تغيير. ولا يعني الحذفَ أبدا."); r += 1
    line(r, "لحذف قيمة", f"اكتب «{SENTINEL}» في العمود الأصفر — صريحةً. وبلا ذلك لا يُفرَّق بين حقلٍ لم تمسّه وحقلٍ أردتَ محوَه.", warn=True); r += 1
    line(r, "القوائم", "ما كان قائمةً (المخرجات، المهارات، المعايير) يُكتب سطرا لكلِّ عنصر داخل الخليّة — بـ Alt+Enter."); r += 2

    line(r, "مثال", "كيف يبدو صفٌّ عُدِّل", head=True); r += 1
    line(r, "اسم الدورة (رمادي)", "دورة التخطيط للمسار المهني وبحث سوق العمل"); r += 1
    line(r, "⟵ تعديلك · اسم الدورة", "دورة بناء المسار المهني  ← يُطبَّق"); r += 1
    line(r, "⟵ تعديلك · المستوى", "(فارغ)  ← لا يتغيّر المستوى"); r += 1
    line(r, "⟵ تعديلك · الاسم القديم", f"{SENTINEL}  ← يُفرَّغ الحقل"); r += 2

    line(r, "الأوراق", "ماذا في كلٍّ منها", head=True); r += 1
    for nm, desc in [
        ("الدورات", "٨١ دورةً بكلِّ حقولها — الاسمُ والإنجليزيُّ والساعاتُ والسعرُ والمخرجاتُ والمهاراتُ والتقييم"),
        ("المحاور", "٤٠٤ محاورَ — العنوانُ والمخرجُ والنشاطُ والأثرُ والساعات"),
        ("محتوى_الدروس", "نصوصُ الدروسِ المكتوبةِ كاملةً بصيغة Markdown"),
        ("المهارات", "كلُّ المهاراتِ ومعها الدوراتُ التي تخدمها"),
        ("مهارات_بلا_دورة", "المهاراتُ المعطّلةُ — وهي ما طلبتَه لتُبنى لها دورات"),
        ("دورات_مقترحة", "تجميعي المقترحُ لتلك المهاراتِ في دورات — اقتراحٌ لا قرار"),
        ("المسارات", "المساراتُ للمرجعِ عند نقلِ دورةٍ من مسارٍ إلى آخر"),
    ]:
        line(r, nm, desc); r += 1
    r += 1

    # أرقامٌ مكتوبةٌ لا صيغ: الملفُّ يُولَّد من المستودعِ في كلِّ مرّة، فالعددُ
    # واقعةُ تصديرٍ لا حسابٌ يتغيّر بتغيّرِ مدخل. وصيغةٌ بلا قيمةٍ مخزّنةٍ تظهر
    # صفرا في كلِّ عارضٍ لا يُعيد الحساب، وهو أسوأُ من رقمٍ صريحٍ صحيح.
    line(r, "الأرقام", "كما صُدّرت من المستودع", head=True); r += 1
    for lbl, val in [
        ("الدورات", n_courses),
        ("المحاور", n_modules),
        ("دروس مكتوبة", f"{n_content} من {n_modules}"),
        ("المهارات", f"{n_skills}  (سجلّ {len(registry)} + امتدادات {len(extensions)})"),
        ("مهارات بلا دورة", n_orphans),
        ("دورات مقترحة", n_prop),
    ]:
        g.cell(row=r, column=1, value=lbl).font = Font(name=FONT, size=10, bold=True)
        cc = g.cell(row=r, column=2, value=val)
        cc.font = Font(name=FONT, size=10)
        cc.alignment = Alignment(horizontal="right")
        g.row_dimensions[r].height = 20
        r += 1
    r += 1

    line(r, "حدودٌ تردّها البوّابة", "قيودٌ في المستودعِ تُرفض عندها القيمةُ آليّا", head=True); r += 1
    line(r, "ساعات الدورة", "عددٌ صحيحٌ بين ١ و٤٠"); r += 1
    line(r, "سعر القائمة", "عددٌ صحيحٌ بين ١٠٠ و٢٠٠، وبعملةٍ واحدةٍ للكتالوجِ كلِّه"); r += 1
    line(r, "المهارات", "كلُّ معرّفٍ لفظيٍّ يجب أن يكون مسجّلا — وإلّا رُفض الاستيراد"); r += 1
    line(r, "لو خالفتَ حدّا", "لا يضيع تعديلُك: أُبلغك بما رُفض ولماذا، وأنتظر قرارَك — إمّا تعديلُ القيمةِ أو تعديلُ الحدِّ نفسِه.", warn=True); r += 2

    line(r, "عند الإعادة", "أرجع الملفَّ كما هو بأوراقِه وأعمدتِه. أقرأ الأصفرَ وحدَه، وأعرض عليك قائمةَ الفروقِ قبل أن أكتب أيَّ شيء."); r += 1

    g.freeze_panes = "A2"

    wb.save(out)
    size = os.path.getsize(out) / (1024 * 1024)
    print(f"✓ {out}  ({size:.1f} MB)")
    print(f"  الدورات {n_courses} · المحاور {n_modules} · دروس مكتوبة {n_content}")
    print(f"  المهارات {n_skills} · بلا دورة {n_orphans} · دورات مقترحة {n_prop}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
