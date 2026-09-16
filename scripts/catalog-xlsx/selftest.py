#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
حارسُ قارئِ الكرّاسة — ورقةُ «دورات_مقترحة» تُقرأ كما تُقرأ أخواتُها.

── لمَ وُجد هذا الحارس ──

كانت `load_sources` تُنشئ `"proposed": {}` فارغةً ولا تقرأ
`proposed-courses.json` أصلا. فكلُّ تعديلٍ في الأعمدةِ الصفراءِ لتلك الورقة —
عنوانا عربيّا كان أو مصطلحا إنجليزيّا أو مستوًى أو ساعاتٍ أو مهاراتٍ — كان
يرتدُّ «معرّفٌ لا أصلَ له في المصدر». وأسوأُ من ارتدادِه أنّ وجودَ مردودٍ
واحدٍ يُخرج القارئَ بـ`1` فيوقف الاستيرادَ كلَّه: فتضيع معه تعديلاتُ ورقةِ
الدوراتِ المقبولةُ التي لا علّةَ فيها.

وكان الصمتُ تامّا في الاتّجاهِ الآخر: `_decision` و`_notes` حقلانِ مِن
`META_FIELDS`، وهما يُفحصان **قبل** `rec is None`، فيمرّان. فيرى المراجعُ
قرارَه مقبولا وعنوانَه مردودا في التقريرِ نفسِه، ولا شيءَ يقول لمَ افترقا.

── والفحصُ على البنية لا على النصّ ──

لا يطابق هذا الحارسُ حرفا في ملفّ: يبني كرّاسةً بالشكلِ الذي يتوقّعه القارئُ
(صفُّ العلامات، صفُّ المفاتيح، ثمّ الصفوف)، ويُشغّل القارئَ نفسَه، ويقرأ
تقريرَه JSON. فلو عاد `"proposed": {}` غدا سقط هذا الفحصُ من أوّل دعوى.

    python3 scripts/catalog-xlsx/selftest.py
"""

from __future__ import annotations

import json
import os
import subprocess
import sys
import tempfile

from openpyxl import Workbook

HERE = os.path.dirname(os.path.abspath(__file__))
READER = os.path.join(HERE, "read_workbook.py")
PROPOSED = os.path.join(HERE, "proposed-courses.json")

EDIT_MARK = "⟵"

# أعمدةُ ورقةِ «دورات_مقترحة» كما يبنيها build_workbook.py: رماديٌّ ثمّ أصفر
COLUMNS = [
    ("proposed_id", False), ("title_ar", True), ("title_term_en", True),
    ("family_id", False), ("level_ar", True), ("suggested_total_hours", True),
    ("skill_slugs", True), ("_decision", True), ("_notes", True),
]

failures: list[str] = []


def check(name: str, ok: bool, detail: str = "") -> None:
    print(f"  {'✓' if ok else '✗'} {name}" + (f"\n      {detail}" if not ok and detail else ""))
    if not ok:
        failures.append(name)


def build(rows: list[dict]) -> str:
    """كرّاسةٌ بورقةِ «دورات_مقترحة» وحدَها، بالشكل الذي يقرؤه القارئ."""
    wb = Workbook()
    ws = wb.active
    ws.title = "دورات_مقترحة"
    ws.cell(1, 1).value = "عنوانُ الورقة — لا يُقرأ"
    col = 1
    pos: dict[str, int] = {}
    for field, editable in COLUMNS:
        ws.cell(2, col).value = field
        ws.cell(3, col).value = field
        if editable:
            col += 1
            ws.cell(2, col).value = f"{EDIT_MARK} تعديلك · {field}"
            ws.cell(3, col).value = "اتركه فارغا = لا تغيير"
            pos[field] = col
        col += 1
    for i, row in enumerate(rows):
        r = 4 + i
        ws.cell(r, 1).value = row.get("proposed_id")
        for field, value in row.items():
            if field != "proposed_id":
                ws.cell(r, pos[field]).value = value
    path = os.path.join(tempfile.mkdtemp(), "كرّاسة.xlsx")
    wb.save(path)
    return path


def run(path: str) -> dict:
    report = path + ".json"
    subprocess.run([sys.executable, READER, path, "--json", report],
                   capture_output=True, check=False)
    with open(report, encoding="utf-8") as f:
        return json.load(f)


def main() -> int:
    with open(PROPOSED, encoding="utf-8") as f:
        proposed = json.load(f)["courses"]
    first = proposed[0]
    pid = first["proposed_id"]

    print("حارسُ قارئِ الكرّاسة · ورقةُ دورات_مقترحة\n")

    # ١ · عنوانٌ جديدٌ لدورةٍ مقترحةٍ يُقبل ولا يرتدّ
    rep = run(build([{"proposed_id": pid, "title_ar": "دورة اسمٌ جديدٌ للفحص"}]))
    hit = [c for c in rep["changes"] if c["id"] == pid and c["field"] == "title_ar"]
    check("تعديلُ العنوانِ العربيِّ يُقبل", bool(hit),
          f"رُدّ بـ: {[p.get('reason') for p in rep['problems']]}")
    check("ولا مردودَ في التقرير", not rep["problems"],
          f"{len(rep['problems'])} مردود")
    if hit:
        check("ويُقابَل بالقيمةِ من proposed-courses.json لا بفراغ",
              hit[0]["from"] == first["title_ar"],
              f"قرأ «{hit[0]['from']}» والمصدرُ «{first['title_ar']}»")

    # ٢ · المطابقُ للمصدرِ لا أثرَ له — دليلُ أنّ المقابلةَ تقع فعلا
    rep = run(build([{"proposed_id": pid, "title_ar": first["title_ar"]}]))
    check("وكتابةُ القيمةِ نفسِها بلا أثر", rep["noops"] == 1 and not rep["changes"],
          f"noops={rep['noops']} changes={len(rep['changes'])}")

    # ٣ · القائمةُ تُقابَل مسرودةً سطرا سطرا لا مصفوفةً
    slugs = "\n".join(first["skill_slugs"])
    rep = run(build([{"proposed_id": pid, "skill_slugs": slugs}]))
    check("والمهاراتُ تُقابَل سطرا سطرا", rep["noops"] == 1,
          f"noops={rep['noops']} changes={len(rep['changes'])}")

    # ٤ · صفٌّ بلا معرّفٍ يبقى اقتراحا جديدا
    rep = run(build([{"proposed_id": None, "title_ar": "دورةٌ مقترحةٌ تماما"}]))
    check("وصفٌّ بلا معرّفٍ يبقى صفّا جديدا", len(rep["new_rows"]) == 1,
          f"new_rows={len(rep['new_rows'])}")

    # ٥ · ولا يصير القارئُ متساهلا: معرّفٌ مخترَعٌ يُردّ كما كان
    rep = run(build([{"proposed_id": "C-NOPE-999", "title_ar": "معرّفٌ مخترَع"}]))
    check("ومعرّفٌ مخترَعٌ يُردّ", len(rep["problems"]) == 1 and not rep["changes"],
          f"problems={len(rep['problems'])} changes={len(rep['changes'])}")

    print()
    if failures:
        print(f"✗ سقط {len(failures)} من الفحوص: " + " · ".join(failures))
        return 1
    print("✓ كلُّ الفحوصِ خضراء")
    return 0


if __name__ == "__main__":
    sys.exit(main())
